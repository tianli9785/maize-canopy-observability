from pathlib import Path
import hashlib
import os
import tempfile
from openpyxl import Workbook, load_workbook
import numpy as np

FIXED_TRAIN_N = 126
FIXED_VAL_N = 54


def _first_data_row_index(rows):
    for i, row in enumerate(rows[1:], start=1):
        if not row:
            continue
        value = row[0]
        try:
            float(value)
            return i
        except (TypeError, ValueError):
            continue
    return 1


def _paired_sheets(sheet_names):
    names = set(sheet_names)
    pairs = []
    for train_name in sheet_names:
        if "-Training" not in train_name:
            continue
        val_name = train_name.replace("-Training", "-Verification", 1)
        if val_name not in names:
            continue
        output_name = train_name.replace("-Training", "", 1)
        pairs.append((train_name, val_name, output_name))
    return pairs


def _combine_workbook(source_path: Path, output_path: Path) -> None:
    src = load_workbook(source_path, data_only=True, read_only=True)
    pairs = _paired_sheets(src.sheetnames)
    if not pairs:
        raise ValueError(f"No Training/Verification sheet pairs found in {source_path}")

    dst = Workbook(write_only=True)
    default = dst.create_sheet("Temporary")

    for train_name, val_name, output_name in pairs:
        train_rows = [tuple(r) for r in src[train_name].iter_rows(values_only=True)]
        val_rows = [tuple(r) for r in src[val_name].iter_rows(values_only=True)]
        val_data_start = _first_data_row_index(val_rows)

        ws = dst.create_sheet(output_name)
        for row in train_rows:
            ws.append(row)
        for row in val_rows[val_data_start:]:
            ws.append(row)

    dst.remove(default)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    dst.save(output_path)
    src.close()


def prepare_fixed_inputs(script_file: str, required_files) -> str:
    script_dir = Path(script_file).resolve().parent
    source_dir = script_dir / "Data"
    cache_key = hashlib.sha1(os.fspath(script_dir).encode("utf-8")).hexdigest()[:12]
    output_dir = Path(tempfile.gettempdir()) / f"target_fixed_inputs_{cache_key}"

    for filename in required_files:
        source_path = source_dir / filename
        if not source_path.exists():
            raise FileNotFoundError(source_path)
        output_path = output_dir / filename
        if (not output_path.exists()) or output_path.stat().st_mtime < source_path.stat().st_mtime:
            _combine_workbook(source_path, output_path)

    return os.fspath(output_dir)


def fixed_indices(n_samples: int):
    expected = FIXED_TRAIN_N + FIXED_VAL_N
    if n_samples != expected:
        raise ValueError(
            f"Expected {expected} samples ({FIXED_TRAIN_N} training + {FIXED_VAL_N} validation), "
            f"but read {n_samples}."
        )
    return np.arange(FIXED_TRAIN_N, dtype=int), np.arange(FIXED_TRAIN_N, expected, dtype=int)
